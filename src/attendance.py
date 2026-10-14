"""
Keep data/<prefix>_attendance.xlsx and its .csv copy in step (SPEC.md).

You may edit either file: in Excel/Numbers (usually the xlsx), by hand (sometimes the csv), or in grade-tui (both).
Each sync saves what it wrote as a baseline (history/attendance_sync.json) and compares each file with it:

    changed in one file only       → copied to the other (cells, Name, session columns added or removed)
    changed differently in both    → a conflict: you choose xlsx / csv / skip
    no baseline yet (first sync)   → a column in one file only counts as deleted in the other, and a differing cell
                                     takes the newer file's value; everything is confirmed first

Removing a column that still holds marks is always asked. Rows follow the roster (course_info/*_student_info.csv):
a roster student missing from a file is added; a row not on the roster is reported and left alone.

The xlsx is edited in place (openpyxl): the cells that differ, deleted or inserted columns, then the total header
`total (<sessions>pts)`, each row's COUNTIF range, the column widths, the P/A/L/X dropdown and the colour rules.
The csv's total uses the same rule as the grade: P = 1, L = 0.8, X = 1, EA = 1, A = 0.
"""

from __future__ import annotations

import hashlib
import json
import os
import re
from copy import copy
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from src.csvio import CsvText, atomic_write_bytes, read_csv_text, render_csv_text

CODE_VALUE = {"P": 1.0, "L": 0.8, "X": 1.0, "EA": 1.0, "A": 0.0}
CODES = ("P", "A", "L", "X")
BASELINE = "attendance_sync.json"
NAME = "Name"


class Skip:
    """A conflict left as it is: each file keeps its own value."""

    def __init__(self, x: str, c: str):
        self.x, self.c = x, c

    def __repr__(self):
        return f"Skip({self.x!r}, {self.c!r})"


def side_value(v, side: str) -> str:
    if isinstance(v, Skip):
        return v.x if side == "xlsx" else v.c
    return v


# ---- reading --------------------------------------------------------------------------------------------------
def norm(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        if v != v:
            return ""
        return str(int(v)) if v.is_integer() else repr(v)
    s = str(v).strip()
    return "" if s.lower() == "nan" else s


def is_total(h: str) -> bool:
    return h.strip().lower().startswith("total")


def is_session(h: str) -> bool:
    h = h.strip()
    return bool(h) and h not in ("Student ID", NAME) and not is_total(h) and not h.startswith("=") \
        and "unnamed" not in h.lower()


@dataclass
class Sheet:
    sessions: list
    rows: dict                                   # sid → {"Name": …, session: code}
    order: list                                  # sids in file order

    def marks(self, session: str) -> int:
        return sum(1 for r in self.rows.values() if r.get(session, ""))


def _sheet(header: list, body) -> Sheet:
    idx = {h.strip(): i for i, h in enumerate(header) if h is not None}
    sessions = [h.strip() for h in header if h is not None and is_session(h)]
    sid_i, name_i = idx.get("Student ID"), idx.get(NAME)
    rows, order = {}, []
    if sid_i is None:
        return Sheet(sessions, rows, order)
    for r in body:
        sid = norm(r[sid_i]) if sid_i < len(r) else ""
        if not sid or sid in rows:
            continue
        rec = {NAME: norm(r[name_i]) if name_i is not None and name_i < len(r) else ""}
        for s in sessions:
            i = idx[s]
            rec[s] = norm(r[i]) if i < len(r) else ""
        rows[sid] = rec
        order.append(sid)
    return Sheet(sessions, rows, order)


def read_csv_sheet(path) -> Sheet:
    doc = read_csv_text(path)
    return _sheet(doc.header, doc.rows)


def _open_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path)
    return wb, wb.worksheets[0]


def read_xlsx_sheet(path) -> Sheet:
    _, ws = _open_xlsx(path)
    header = [None if c.value is None else str(c.value) for c in ws[1]]
    body = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
    return _sheet(header, body)


# ---- files ----------------------------------------------------------------------------------------------------
def files_of(course) -> tuple[Path | None, Path | None]:
    """(xlsx, csv) of the course's attendance, either may be None."""
    data = Path(course) / "data"
    if not data.is_dir():
        return None, None
    xs = sorted(p for p in data.glob("*attendance.xlsx") if not p.name.startswith(("~$", ".")))
    cs = sorted(p for p in data.glob("*attendance.csv") if not p.name.startswith("."))
    pref = [p for p in xs if p.name.endswith("_attendance.xlsx")]
    x = (pref or xs or [None])[0]
    if x is not None:
        return x, x.with_suffix(".csv")
    pref = [p for p in cs if p.name.endswith("_attendance.csv")]
    return None, (pref or cs or [None])[0]


def locked(xlsx: Path) -> bool:
    return (xlsx.parent / f"~${xlsx.name}").exists()


def sha(path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest() if path and Path(path).exists() else ""


def roster(course) -> dict | None:
    """{sid: name} from course_info/*_student_info.csv, or None when there is no roster."""
    info = Path(course) / "course_info"
    files = sorted(info.glob("*_student_info.csv")) if info.is_dir() else []
    if not files:
        return None
    doc = read_csv_text(files[0])
    si, ni = doc.col("Student ID"), doc.col(NAME)
    if si is None:
        return None
    out = {}
    for r in doc.rows:
        sid = norm(r[si])
        if sid:
            out[sid] = norm(r[ni]) if ni is not None else ""
    return out


def baseline_path(course) -> Path:
    return Path(course) / "history" / BASELINE


def load_baseline(course) -> dict | None:
    p = baseline_path(course)
    if not p.exists():
        return None
    try:
        b = json.loads(p.read_text(encoding="utf-8"))
        return b if isinstance(b, dict) and "rows" in b else None
    except ValueError:
        return None


def save_baseline(course, sessions: list, rows: dict, xlsx, csv) -> None:
    p = baseline_path(course)
    p.parent.mkdir(exist_ok=True)
    data = {"ts": datetime.now().isoformat(timespec="seconds"), "sessions": sessions, "rows": rows,
            "xlsx": xlsx.name if xlsx else None, "sha": {"xlsx": sha(xlsx), "csv": sha(csv)}}
    atomic_write_bytes(p, json.dumps(data, ensure_ascii=False, indent=1).encode("utf-8"))


# ---- totals ---------------------------------------------------------------------------------------------------
def att_total(codes) -> float:
    return sum(CODE_VALUE.get(str(c).strip(), 0.0) for c in codes)


def fmt_total(x: float) -> str:
    return str(int(x)) if float(x).is_integer() else f"{x:.2f}".rstrip("0").rstrip(".")


def total_header(n: int) -> str:
    return f"total ({n}pts)"


def _same_num(a: str, b: str) -> bool:
    try:
        return abs(float(a) - float(b)) < 1e-9
    except (TypeError, ValueError):
        return False


def session_date(h: str) -> date | None:
    try:
        return datetime.strptime(h.strip(), "%d %b %Y").date()
    except ValueError:
        return None


# ---- plan -----------------------------------------------------------------------------------------------------
@dataclass
class Removal:
    session: str
    deleted_in: str          # "xlsx" / "csv": the file it is gone from
    marks: int               # marks it holds in the other file


@dataclass
class AttPlan:
    course: Path
    xlsx: Path
    csv: Path
    X: Sheet
    C: Sheet | None          # None: no csv yet (it is made from the xlsx)
    B: dict | None
    first_run: bool
    newer: str
    order: list = field(default_factory=list)          # every session of either file, merged order
    removals: list = field(default_factory=list)       # Removal
    additions: list = field(default_factory=list)      # (session, added_in)
    cells: dict = field(default_factory=dict)          # (sid, field) → value or Skip, for rows in both files
    copies: list = field(default_factory=list)         # (sid, field, x, c, to)
    conflicts: list = field(default_factory=list)      # (sid, field, x, c, base)
    rows_added: list = field(default_factory=list)     # (sid, to: "xlsx"/"csv"/"both")
    orphans: list = field(default_factory=list)        # (sid, name, where)
    empty_past: list = field(default_factory=list)
    stale_total: list = field(default_factory=list)    # file names whose total column needs refreshing
    is_locked: bool = False

    @property
    def marked(self) -> list:
        return [r for r in self.removals if r.marks]

    def needs_question(self) -> bool:
        return bool(self.conflicts or self.marked or (self.first_run and (self.copies or self.removals
                                                                            or self.additions)))

    def has_changes(self) -> bool:
        return bool(self.removals or self.additions or self.copies or self.conflicts or self.rows_added
                    or self.stale_total or self.C is None)


def _merge_order(a: list, b: list) -> list:
    out = list(a)
    for i, s in enumerate(b):
        if s in out:
            continue
        prev = next((b[j] for j in range(i - 1, -1, -1) if b[j] in out), None)
        out.insert(out.index(prev) + 1 if prev else 0, s)
    return out


def plan(course, today: date | None = None) -> AttPlan | None:
    """What a sync would do, or None when the course has no attendance xlsx."""
    course = Path(course)
    xlsx, csv = files_of(course)
    if xlsx is None:
        return None
    today = today or date.today()
    X = read_xlsx_sheet(xlsx)
    C = read_csv_sheet(csv) if csv.exists() else None
    B = load_baseline(course)
    if B is not None and B.get("xlsx") not in (None, xlsx.name):
        B = None
    newer = "csv" if C is not None and csv.stat().st_mtime > xlsx.stat().st_mtime else "xlsx"
    p = AttPlan(course, xlsx, csv, X, C, B, first_run=B is None and C is not None, newer=newer,
                is_locked=locked(xlsx))
    if C is None:
        p.order = list(X.sessions)
        p.stale_total = [csv.name]
        return p

    bs = set(B["sessions"]) if B else set()
    brows = B["rows"] if B else {}
    p.order = _merge_order(X.sessions, C.sessions)
    for s in p.order:
        inx, inc = s in X.sessions, s in C.sessions
        if inx and inc:
            continue
        if B is None or s in bs:
            gone = "csv" if inx else "xlsx"
            p.removals.append(Removal(s, gone, (X if inx else C).marks(s)))
        else:
            p.additions.append((s, "xlsx" if inx else "csv"))

    # rows
    ros = roster(course)
    for sid in _merge_order(X.order, C.order):
        inx, inc = sid in X.rows, sid in C.rows
        if inx and inc:
            if ros is not None and sid not in ros:
                p.orphans.append((sid, X.rows[sid][NAME], "both files, not on the roster"))
            continue
        name = (X.rows.get(sid) or C.rows.get(sid))[NAME]
        if ros is None or sid in ros:
            p.rows_added.append((sid, "csv" if inx else "xlsx"))
        else:
            p.orphans.append((sid, name, f"only in the {'xlsx' if inx else 'csv'}, not on the roster"))
    for sid in (ros or {}):
        if sid not in X.rows and sid not in C.rows:
            p.rows_added.append((sid, "both"))

    # cells of rows in both files, in the columns both files have
    for sid in X.order:
        if sid not in C.rows:
            continue
        for f in [NAME] + [s for s in p.order if s in X.sessions and s in C.sessions]:
            x, c = X.rows[sid].get(f, ""), C.rows[sid].get(f, "")
            if x == c:
                p.cells[(sid, f)] = x
                continue
            b = brows.get(sid, {}).get(f) if f in bs or f == NAME else None
            if b is not None and c == b:
                p.cells[(sid, f)] = x
                p.copies.append((sid, f, x, c, "csv"))
            elif b is not None and x == b:
                p.cells[(sid, f)] = c
                p.copies.append((sid, f, x, c, "xlsx"))
            elif B is None:
                v = x if newer == "xlsx" else c
                p.cells[(sid, f)] = v
                p.copies.append((sid, f, x, c, "csv" if newer == "xlsx" else "xlsx"))
            else:
                p.cells[(sid, f)] = Skip(x, c)
                p.conflicts.append((sid, f, x, c, b))

    # empty sessions with a date before today (offered, never automatic)
    removed = {r.session for r in p.removals}
    for s in p.order:
        if s in removed:
            continue
        d = session_date(s)
        if d is None or d >= today:
            continue
        if not (X.marks(s) if s in X.sessions else 0) and not (C.marks(s) if s in C.sessions else 0):
            p.empty_past.append(s)

    # totals
    doc = read_csv_text(csv)
    ti = next((i for i, h in enumerate(doc.header) if is_total(h)), None)
    if ti is None or doc.header[ti] != total_header(len(C.sessions)) or _csv_totals_stale(doc, ti, C):
        p.stale_total.append(csv.name)
    if not _xlsx_total_ok(xlsx, X):
        p.stale_total.append(xlsx.name)
    return p


def _csv_totals_stale(doc: CsvText, ti: int, C: Sheet) -> bool:
    si = doc.col("Student ID")
    for r in doc.rows:
        sid = norm(r[si]) if si is not None else ""
        if sid in C.rows:
            want = fmt_total(att_total(C.rows[sid].get(s, "") for s in C.sessions))
            if not _same_num(r[ti] if ti < len(r) else "", want):
                return True
    return False


def _xlsx_total_ok(xlsx: Path, X: Sheet) -> bool:
    _, ws = _open_xlsx(xlsx)
    cols = _xlsx_cols(ws)
    if cols["total"] is None:
        return False
    if str(ws.cell(row=1, column=cols["total"]).value).strip() != total_header(len(X.sessions)):
        return False
    rng = _session_range(cols)
    for r in _data_rows(ws, cols):
        v = ws.cell(row=r, column=cols["total"]).value
        if not (isinstance(v, str) and v.startswith("=")) or _fix_formula(v, rng, r) != v:
            return False
    return True


# ---- decisions → the target ------------------------------------------------------------------------------------
@dataclass
class Decisions:
    keep: set = field(default_factory=set)          # sessions with marks whose removal you refused
    drop_empty: bool = False                        # remove the empty past sessions
    choose: dict = field(default_factory=dict)      # (sid, field) → "x" / "c" for conflicts (else skipped)
    apply: bool = True


@dataclass
class Target:
    sessions: list
    values: dict             # sid → {field: value or Skip} for rows written in both files
    new_rows: dict           # side → [sid] to append
    baseline_rows: dict      # sid → {field: value}


def target(p: AttPlan, d: Decisions) -> Target:
    drop = {r.session for r in p.removals if r.session not in d.keep}
    if d.drop_empty:
        drop |= set(p.empty_past)
    sessions = [s for s in p.order if s not in drop]
    X, C = p.X, p.C
    ros = roster(p.course) or {}
    values, base = {}, {}
    if C is None:
        for sid in X.order:
            values[sid] = dict(X.rows[sid])
        return Target(sessions, values, {"xlsx": [], "csv": list(X.order)}, values)
    new_rows = {"xlsx": [], "csv": []}
    added = {sid: to for sid, to in p.rows_added}
    for sid in _merge_order(X.order, C.order) + [s for s, to in p.rows_added if to == "both"]:
        if sid in X.rows and sid in C.rows:
            rec = {}
            for f in [NAME] + sessions:
                if (sid, f) in p.cells:
                    v = p.cells[(sid, f)]
                    if isinstance(v, Skip) and (sid, f) in d.choose:
                        v = v.x if d.choose[(sid, f)] == "x" else v.c
                    rec[f] = v
                else:  # a column only one file has (added, or a removal you refused)
                    rec[f] = X.rows[sid].get(f, "") if f in X.sessions else C.rows[sid].get(f, "")
            values[sid] = rec
        elif sid in added:
            to = added[sid]
            src = {NAME: ros.get(sid, "")} if to == "both" else (X if to == "csv" else C).rows[sid]
            values[sid] = {f: src.get(f, "") for f in [NAME] + sessions}
            for side in (["xlsx", "csv"] if to == "both" else [to]):
                new_rows[side].append(sid)
        else:
            continue
        base[sid] = {f: v for f, v in values[sid].items() if not isinstance(v, Skip)}
    return Target(sessions, values, new_rows, base)


# ---- writing --------------------------------------------------------------------------------------------------
def render_csv(path: Path, t: Target, drop_rows: set = frozenset(), keep_columns: bool = False) -> bytes:
    """The csv bytes for target `t` (rows not in t.values keep their cells; columns follow t.sessions)."""
    if path.exists():
        doc = read_csv_text(path)
    else:
        doc = CsvText(header=["Student ID", NAME], rows=[], bom=False, eol="\n")
    idx = {h.strip(): i for i, h in enumerate(doc.header)}
    si, ni = idx.get("Student ID"), idx.get(NAME)
    ti = next((i for i, h in enumerate(doc.header) if is_total(h)), None)
    sessions = [h.strip() for h in doc.header if is_session(h)] if keep_columns else t.sessions
    extra = [h for h in doc.header if h.strip() not in ("Student ID", NAME) and not is_total(h)
             and not is_session(h)]
    header = ["Student ID", NAME] + sessions + extra + [total_header(len(sessions))]
    rows = []

    def build(sid, name, cells, old_total):
        tot = fmt_total(att_total(cells.get(s, "") for s in sessions))
        if _same_num(old_total, tot):
            tot = old_total
        return [sid, name] + [cells.get(s, "") for s in sessions] + [cells.get(h, "") for h in extra] + [tot]

    for r in doc.rows:
        get = (lambda i: r[i] if i is not None and i < len(r) else "")
        sid = norm(get(si))
        if sid and sid in drop_rows:
            continue
        old = {h.strip(): r[i] for i, h in enumerate(doc.header) if i < len(r)}
        name = get(ni)
        if sid and sid in t.values:
            v = t.values[sid]
            for f in sessions:
                if f in v:
                    nv = side_value(v[f], "csv")
                    if norm(old.get(f, "")) != nv:
                        old[f] = nv
            nv = side_value(v.get(NAME, norm(name)), "csv")
            if norm(name) != nv:
                name = nv
        rows.append(build(get(si), name, old, get(ti)) if sid else r)
    for sid in t.new_rows.get("csv", []):
        v = t.values[sid]
        rows.append(build(sid, side_value(v.get(NAME, ""), "csv"),
                          {f: side_value(v.get(f, ""), "csv") for f in sessions}, ""))
    doc.header, doc.rows = header, rows
    return render_csv_text(doc)


def _xlsx_cols(ws) -> dict:
    out = {"sid": None, "name": None, "total": None, "sessions": {}}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=1, column=c).value
        if v is None:
            continue
        h = str(v).strip()
        if h == "Student ID":
            out["sid"] = c
        elif h == NAME:
            out["name"] = c
        elif is_total(h):
            out["total"] = c
        elif is_session(h):
            out["sessions"][h] = c
    return out


def _data_rows(ws, cols) -> list:
    if cols["sid"] is None:
        return []
    return [r for r in range(2, ws.max_row + 1) if norm(ws.cell(row=r, column=cols["sid"]).value)]


def _session_range(cols) -> tuple | None:
    idx = sorted(cols["sessions"].values())
    if not idx:
        return None
    from openpyxl.utils import get_column_letter
    return get_column_letter(idx[0]), get_column_letter(idx[-1])


_RANGE = re.compile(r"\$?[A-Z]{1,3}\$?\d+:\$?[A-Z]{1,3}\$?\d+")


def _fix_formula(f: str, rng, row: int) -> str:
    if rng is None:
        return "=0"
    new = f"{rng[0]}{row}:{rng[1]}{row}"
    return _RANGE.sub(new, f) if _RANGE.search(f) else _standard_formula(rng, row)


def _standard_formula(rng, row: int) -> str:
    if rng is None:
        return "=0"
    r = f"{rng[0]}{row}:{rng[1]}{row}"
    return f'=COUNTIF({r}, "P") + COUNTIF({r}, "L")*0.8 + COUNTIF({r}, "X")'


def _copy_style(src, dst):
    if src.has_style:
        dst._style = copy(src._style)


def edit_xlsx(wb, ws, t: Target, drop_rows: set = frozenset(), keep_columns: bool = False) -> int:
    """Applies target `t` to the worksheet in place; returns the number of changes (0 = leave the file alone)."""
    from openpyxl.utils import get_column_letter
    changes = 0
    cols = _xlsx_cols(ws)
    widths = {}
    for c in range(1, ws.max_column + 1):
        h = ws.cell(row=1, column=c).value
        dim = ws.column_dimensions.get(get_column_letter(c))
        if h is not None and dim is not None and dim.width:
            widths[str(h).strip()] = dim.width
    session_width = next((widths[s] for s in cols["sessions"] if s in widths), 6)

    # rows to drop (bottom up)
    for r in sorted((r for r in _data_rows(ws, cols) if norm(ws.cell(row=r, column=cols["sid"]).value)
                     in drop_rows), reverse=True):
        ws.delete_rows(r)
        changes += 1

    # columns: remove, then insert, so the sessions run in t.sessions order
    if not keep_columns:
        for s, c in sorted(cols["sessions"].items(), key=lambda kv: -kv[1]):
            if s not in t.sessions:
                ws.delete_cols(c)
                changes += 1
        cols = _xlsx_cols(ws)
        for i, s in enumerate(t.sessions):
            if s in cols["sessions"]:
                continue
            prev = next((t.sessions[j] for j in range(i - 1, -1, -1) if t.sessions[j] in cols["sessions"]), None)
            if prev is not None:
                at = cols["sessions"][prev] + 1
            elif cols["sessions"]:
                at = min(cols["sessions"].values())
            else:
                at = (cols["name"] or cols["sid"] or 1) + 1
            neighbour = at - 1 if prev is not None else at
            ws.insert_cols(at)
            src_col = neighbour if neighbour < at else at + 1
            for r in range(1, ws.max_row + 1):
                _copy_style(ws.cell(row=r, column=src_col), ws.cell(row=r, column=at))
            ws.cell(row=1, column=at).value = s
            changes += 1
            cols = _xlsx_cols(ws)
    sessions = sorted(cols["sessions"], key=lambda s: cols["sessions"][s])

    # cells of the rows already there
    for r in _data_rows(ws, cols):
        sid = norm(ws.cell(row=r, column=cols["sid"]).value)
        v = t.values.get(sid)
        if v is None:
            continue
        for f in [NAME] + sessions:
            c = cols["name"] if f == NAME else cols["sessions"][f]
            if c is None or f not in v:
                continue
            cell = ws.cell(row=r, column=c)
            nv = side_value(v[f], "xlsx")
            if norm(cell.value) != nv:
                cell.value = nv if nv != "" else None
                changes += 1

    # new rows
    rows = _data_rows(ws, cols)
    last = rows[-1] if rows else 1
    for sid in t.new_rows.get("xlsx", []):
        last += 1
        ws.insert_rows(last) if last <= ws.max_row else None
        for c in range(1, ws.max_column + 1):
            if last > 2:
                _copy_style(ws.cell(row=last - 1, column=c), ws.cell(row=last, column=c))
        ws.cell(row=last, column=cols["sid"]).value = sid
        v = t.values[sid]
        if cols["name"]:
            ws.cell(row=last, column=cols["name"]).value = side_value(v.get(NAME, ""), "xlsx") or None
        for f in sessions:
            nv = side_value(v.get(f, ""), "xlsx")
            ws.cell(row=last, column=cols["sessions"][f]).value = nv or None
        changes += 1

    # total column, formulas, widths, dropdown, colours
    cols = _xlsx_cols(ws)
    if cols["total"] is None:
        cols["total"] = ws.max_column + 1
        _copy_style(ws.cell(row=1, column=cols["total"] - 1), ws.cell(row=1, column=cols["total"]))
    hdr = total_header(len(sessions))
    if str(ws.cell(row=1, column=cols["total"]).value or "").strip() != hdr:
        ws.cell(row=1, column=cols["total"]).value = hdr
        changes += 1
    rng = _session_range(cols)
    rows = _data_rows(ws, cols)
    for r in rows:
        cell = ws.cell(row=r, column=cols["total"])
        v = cell.value
        want = _fix_formula(v, rng, r) if isinstance(v, str) and v.startswith("=") else _standard_formula(rng, r)
        if v != want:
            cell.value = want
            changes += 1
    if changes:
        total_width = next((w for k, w in widths.items() if is_total(k)), None)
        for c in range(1, ws.max_column + 1):
            h = ws.cell(row=1, column=c).value
            if h is None:
                continue
            h = str(h).strip()
            w = widths.get(h) or (session_width if h in cols["sessions"] else total_width if is_total(h) else None)
            if w:
                ws.column_dimensions[get_column_letter(c)].width = w
        from openpyxl.utils import column_index_from_string
        for letter in [k for k in ws.column_dimensions if column_index_from_string(k) > ws.max_column]:
            del ws.column_dimensions[letter]  # widths left behind by deleted columns
        _fix_ranges(ws, rng, rows[-1] if rows else 1)
    return changes


def _fix_ranges(ws, rng, last_row: int) -> None:
    from openpyxl.formatting.formatting import ConditionalFormattingList
    area = f"{rng[0]}2:{rng[1]}{last_row}" if rng and last_row >= 2 else None
    for dv in list(ws.data_validations.dataValidation):
        if dv.formula1 and dv.formula1.replace(" ", "") == '"P,A,L,X"':
            if area:
                dv.sqref = area
            else:
                ws.data_validations.dataValidation.remove(dv)
    new = ConditionalFormattingList()
    for cf in ws.conditional_formatting:
        ours = all(getattr(r, "type", None) == "cellIs" and r.formula and r.formula[0].strip('"') in CODES
                   for r in cf.rules)
        where = area if ours else str(cf.sqref)
        if where:
            for rule in cf.rules:
                new.add(where, rule)
    ws.conditional_formatting = new


def save_xlsx(wb, path: Path) -> None:
    tmp = path.with_name(f".{path.stem}.tmp-{os.getpid()}.xlsx")
    try:
        wb.save(tmp)
        if path.exists():
            os.chmod(tmp, path.stat().st_mode & 0o7777)
        os.replace(tmp, path)
    finally:
        if tmp.exists():
            tmp.unlink()


class Locked(Exception):
    pass


def apply(p: AttPlan, d: Decisions, entry) -> dict:
    """Writes both files (only what differs) into history entry `entry`, then saves the baseline.
    Returns {file name: number of changes}."""
    if p.is_locked:
        raise Locked(f"{p.xlsx.name} is open in Excel (~${p.xlsx.name}): close it first")
    t = target(p, d)
    out = {}
    data = render_csv(p.csv, t)
    wb, ws = _open_xlsx(p.xlsx)
    n_x = edit_xlsx(wb, ws, t)
    csv_changed = not p.csv.exists() or p.csv.read_bytes() != data
    if n_x:
        entry.capture(p.xlsx)
    if csv_changed:
        entry.capture(p.csv)
    if n_x or csv_changed:
        entry.capture(baseline_path(p.course))  # so `grader undo` puts the earlier baseline back too
    if n_x:
        save_xlsx(wb, p.xlsx)
        out[p.xlsx.name] = n_x
    if csv_changed:
        atomic_write_bytes(p.csv, data)
        out[p.csv.name] = 1
    save_baseline(p.course, t.sessions, t.baseline_rows, p.xlsx, p.csv)
    return out


# ---- entry points ---------------------------------------------------------------------------------------------
def in_sync_fast(course) -> bool:
    """True when neither file changed since the last sync (no parsing needed)."""
    b = load_baseline(course)
    xlsx, csv = files_of(course)
    if not b or xlsx is None or b.get("xlsx") != xlsx.name:
        return False
    s = b.get("sha") or {}
    return s.get("xlsx") == sha(xlsx) and s.get("csv") == sha(csv)


def describe(p: AttPlan) -> list[str]:
    """Human-readable lines for the preview."""
    out = []
    if p.C is None:
        out.append(f"{p.csv.name} doesn't exist yet: made from {p.xlsx.name}")
        return out
    if p.first_run:
        out.append(f"First attendance sync (no earlier sync to compare with): the newer file, {p.newer}, is proposed "
                   "where they differ.")
    for r in p.removals:
        what = f"{r.marks} mark(s)" if r.marks else "empty"
        other = "csv" if r.deleted_in == "xlsx" else "xlsx"
        out.append(f"  − session '{r.session}': gone from the {r.deleted_in} → remove from the {other} ({what})")
    for s, where in p.additions:
        out.append(f"  + session '{s}': added in the {where} → add to the {'csv' if where == 'xlsx' else 'xlsx'}")
    for sid, f, x, c, to in p.copies:
        new, old = (x, c) if to == "csv" else (c, x)
        out.append(f"  {sid} {f}: {old or '—'} → {new or '—'} in the {to}")
    for sid, to in p.rows_added:
        out.append(f"  + row {sid}: on the roster, missing from the {to if to != 'both' else 'xlsx and csv'}")
    for sid, name, where in p.orphans:
        out.append(f"  ⚠️  {sid} {name}: {where} (left as it is; `grader update` handles the roster)")
    for sid, f, x, c, b in p.conflicts:
        out.append(f"  ≠ {sid} {f}: xlsx {x or '—'}, csv {c or '—'} (last sync {b or '—'})")
    if p.stale_total:
        out.append(f"  total column out of date in: {', '.join(p.stale_total)}")
    return out


def warnings(course) -> list[str]:
    """What needs `grader sync` (for grade-tui); empty when the files are in step."""
    try:
        if in_sync_fast(course):
            return []
        p = plan(course)
    except Exception as e:
        return [f"attendance: could not compare the xlsx and csv ({e})"]
    if p is None or not p.needs_question():
        return []
    bits = []
    if p.first_run:
        bits.append("the xlsx and csv differ and were never synced")
    if p.conflicts:
        bits.append(f"{len(p.conflicts)} cell(s) changed differently in the xlsx and the csv")
    if p.marked:
        bits.append("session(s) with marks removed in one file: " + ", ".join(r.session for r in p.marked))
    return [f"attendance: {'; '.join(bits)} — run `grader sync`"] if bits else []


def auto_sync(course, tool: str = "grade-tui") -> list[str]:
    """On opening a course: applies a sync that needs no question (SPEC decision 6); otherwise returns warnings."""
    from src.history import Entry
    course = Path(course)
    try:
        if in_sync_fast(course):
            return []
        p = plan(course)
        if p is None:
            return []
        if p.needs_question():
            return warnings(course)
        if p.is_locked:
            return [f"attendance: {p.xlsx.name} is open in Excel; not synced"] if p.has_changes() else []
        e = Entry(course, tool, "attendance")
        done = apply(p, Decisions(), e)
        e.close(sum(done.values()) or None)
        return []
    except Exception as ex:
        return [f"attendance: could not sync ({ex})"]


def align_rows(course, entry=None, tool: str = "grader", action: str = "update-attendance",
               class_dates: list | None = None) -> dict:
    """`grader update`: rows follow the roster in both files; the session columns stay as they are.
    Placeholder columns a1…aN (never dated) take `class_dates` by position. Returns {file: changes}."""
    from src.history import Entry
    course = Path(course)
    xlsx, csv = files_of(course)
    ros = roster(course)
    if xlsx is None or ros is None:
        return {}
    if locked(xlsx):
        raise Locked(f"{xlsx.name} is open in Excel (~${xlsx.name}): close it first")
    own = entry is None
    e = entry or Entry(course, tool, action)
    out = {}
    wb, ws = _open_xlsx(xlsx)
    X = read_xlsx_sheet(xlsx)
    C = read_csv_sheet(csv) if csv.exists() else X
    rename = {}
    if class_dates and all(re.fullmatch(r"a\d+", s) for s in X.sessions) and X.sessions \
            and not all(re.fullmatch(r"a\d+", s) for s in class_dates):
        rename = dict(zip(X.sessions, class_dates))

    def side(S: Sheet):
        drop = {sid for sid in S.rows if sid not in ros and not any(S.rows[sid].get(s) for s in S.sessions)}
        new = [sid for sid in ros if sid not in S.rows]
        return drop, new

    for name, S in (("xlsx", X), ("csv", C)):
        if name == "csv" and not csv.exists():
            continue
        drop, new = side(S)
        sessions = [rename.get(s, s) for s in S.sessions]
        values = {sid: {NAME: ros[sid]} for sid in new}
        for sid in new:
            values[sid].update({s: "" for s in sessions})
        t = Target(sessions, values, {name: new}, {})
        if name == "xlsx":
            if rename:
                cols = _xlsx_cols(ws)
                for old, newname in rename.items():
                    ws.cell(row=1, column=cols["sessions"][old]).value = newname
            n = edit_xlsx(wb, ws, t, drop_rows=drop, keep_columns=True) + len(rename)
            if n:
                e.capture(xlsx)
                e.capture(baseline_path(course))
                save_xlsx(wb, xlsx)
                out[xlsx.name] = n
        else:
            data = _render_with_rename(csv, t, drop, rename)
            if csv.read_bytes() != data:
                e.capture(csv)
                e.capture(baseline_path(course))
                atomic_write_bytes(csv, data)
                out[csv.name] = 1
    if not csv.exists():
        apply(plan(course), Decisions(), e)
        out[csv.name] = 1
    if own:
        e.close(sum(out.values()) or None)
    if out:  # both files changed the same way: they stay in step with the last sync
        X2, C2 = read_xlsx_sheet(xlsx), read_csv_sheet(csv)
        if X2.sessions == C2.sessions and X2.rows == C2.rows:
            save_baseline(course, X2.sessions, X2.rows, xlsx, csv)
    return out


def _render_with_rename(csv: Path, t: Target, drop: set, rename: dict) -> bytes:
    if rename and csv.exists():
        doc = read_csv_text(csv)
        doc.header = [rename.get(h.strip(), h) for h in doc.header]
        tmp = csv.with_name(f".{csv.name}.rename-{os.getpid()}")
        try:
            tmp.write_bytes(render_csv_text(doc))
            return render_csv(tmp, t, drop_rows=drop, keep_columns=True)
        finally:
            tmp.unlink(missing_ok=True)
    return render_csv(csv, t, drop_rows=drop, keep_columns=True)


def make_csv(course, entry=None) -> None:
    """After `grader init`/`mock` made the xlsx: write the csv from it and save the first baseline."""
    from src.history import Entry
    p = plan(course)
    if p is None:
        return
    e = entry or Entry(course, "grader", "init")
    apply(p, Decisions(), e)
    if entry is None:
        e.close()
