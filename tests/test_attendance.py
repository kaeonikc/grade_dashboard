import csv
import io
import os
import time
from datetime import date

import openpyxl
import pandas as pd
import pytest

import grader
from src import attendance, tui_api
from src.history import list_entries

SESSIONS = ["1 Sep 2026", "8 Sep 2026", "15 Sep 2026", "22 Sep 2026", "29 Sep 2026"]
TODAY = date(2026, 10, 14)


def xlsx_of(course):
    return next((course / "data").glob("*_attendance.xlsx"))


def csv_rows(course):
    p = xlsx_of(course).with_suffix(".csv")
    return list(csv.reader(io.StringIO(p.read_bytes().decode("utf-8-sig"), newline="")))


def ws_of(course):
    return openpyxl.load_workbook(xlsx_of(course)).worksheets[0]


def header(ws):
    return [c.value for c in ws[1]]


def set_xlsx(course, sid, session, value):
    wb = openpyxl.load_workbook(xlsx_of(course))
    ws = wb.worksheets[0]
    h = header(ws)
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(row=r, column=1).value) == sid:
            ws.cell(row=r, column=h.index(session) + 1).value = value
    wb.save(xlsx_of(course))
    bump(xlsx_of(course))


def delete_xlsx_col(course, session):
    wb = openpyxl.load_workbook(xlsx_of(course))
    ws = wb.worksheets[0]
    ws.delete_cols(header(ws).index(session) + 1)
    wb.save(xlsx_of(course))
    bump(xlsx_of(course))


def edit_csv(course, fn):
    p = xlsx_of(course).with_suffix(".csv")
    rows = csv_rows(course)
    fn(rows)
    buf = io.StringIO()
    csv.writer(buf, lineterminator="\n").writerows(rows)
    p.write_text(buf.getvalue(), encoding="utf-8")
    bump(p)


def bump(p):
    t = time.time() + 5 + bump.n
    bump.n += 1
    os.utime(p, (t, t))


bump.n = 0


@pytest.fixture
def att(course):
    """The test course plus an attendance sheet (5 dated sessions), synced once as grader init does."""
    reg = pd.read_csv(next((course / "course_info").glob("*_student_info.csv")), dtype=str)
    path = course / "data" / "2026_S1_Test_SEC_51_attendance.xlsx"
    grader.save_attendance_excel(path, reg, SESSIONS)
    attendance.make_csv(course)
    for sid, codes in {"69000001": "PPP", "69000002": "PLA", "69000003": "PPX"}.items():
        for s, c in zip(SESSIONS, codes):
            set_xlsx(course, sid, s, c)
    sync(course, yes=True)
    return course


def sync(course, yes=False, answers=None, monkeypatch=None):
    if answers is not None:
        it = iter(answers)
        monkeypatch.setattr(grader, "_ask", lambda q, default="": next(it))
    try:
        attendance.plan.__defaults__ = (TODAY,)
        grader.sync_cmd(str(course), yes=yes)
    finally:
        attendance.plan.__defaults__ = (None,)


def test_init_writes_csv_with_totals_and_a_baseline(att):
    rows = csv_rows(att)
    assert rows[0] == ["Student ID", "Name"] + SESSIONS + ["total (5pts)"]
    assert [r[-1] for r in rows[1:]] == ["3", "1.8", "3"]
    assert attendance.in_sync_fast(att)


def test_xlsx_edit_reaches_csv_on_open_and_csv_edit_reaches_xlsx(att):
    set_xlsx(att, "69000003", "15 Sep 2026", "A")
    assert attendance.warnings(att) == []
    tui_api.get_course_data(att)  # opening the course syncs it (one-sided, no question)
    assert csv_rows(att)[3][4] == "A" and csv_rows(att)[3][-1] == "2"
    assert [e["action"] for e in list_entries(att)][-1] == "attendance"

    edit_csv(att, lambda r: r[1].__setitem__(5, "L"))  # 69000001, 22 Sep
    sync(att, yes=True)
    ws = ws_of(att)
    assert ws.cell(row=2, column=6).value == "L"
    # formatting and the dropdown survive
    assert ws.cell(row=1, column=3).alignment.textRotation == 90
    assert any('"P,A,L,X"' in (dv.formula1 or "") for dv in ws.data_validations.dataValidation)


def test_empty_column_deleted_in_xlsx_is_mirrored_and_totals_follow(att):
    delete_xlsx_col(att, "29 Sep 2026")
    sync(att, yes=True)
    rows = csv_rows(att)
    assert rows[0] == ["Student ID", "Name"] + SESSIONS[:4] + ["total (4pts)"]
    ws = ws_of(att)
    assert header(ws)[-1] == "total (4pts)"
    assert "C2:F2" in ws.cell(row=2, column=7).value
    dv = next(dv for dv in ws.data_validations.dataValidation if '"P,A,L,X"' in (dv.formula1 or ""))
    assert str(dv.sqref) == "C2:F4"
    assert {str(cf.sqref) for cf in ws.conditional_formatting} == {"C2:F4"}
    data = tui_api.get_course_data(att)
    assert len(data["raw_scores"]) == 3 and attendance.in_sync_fast(att)


def test_removing_a_marked_column_is_asked_and_no_restores_it(att, monkeypatch):
    edit_csv(att, lambda rows: [r.pop(2) for r in rows])  # "1 Sep 2026" has marks
    assert "run `grader sync`" in attendance.warnings(att)[0]
    tui_api.get_course_data(att)  # not applied on open
    assert "1 Sep 2026" not in csv_rows(att)[0]
    sync(att, answers=["n", "n"], monkeypatch=monkeypatch)
    assert csv_rows(att)[0][2] == "1 Sep 2026" and csv_rows(att)[1][2] == "P"
    assert header(ws_of(att))[2] == "1 Sep 2026"

    edit_csv(att, lambda rows: [r.pop(2) for r in rows])
    sync(att, yes=True)  # --yes keeps it
    assert csv_rows(att)[0][2] == "1 Sep 2026"
    edit_csv(att, lambda rows: [r.pop(2) for r in rows])
    sync(att, answers=["y", "n"], monkeypatch=monkeypatch)
    assert "1 Sep 2026" not in header(ws_of(att)) and csv_rows(att)[0][-1] == "total (4pts)"


def test_conflict_is_asked_and_yes_skips_it(att, monkeypatch):
    set_xlsx(att, "69000002", "8 Sep 2026", "P")
    edit_csv(att, lambda r: r[2].__setitem__(3, "A"))
    p = attendance.plan(att, TODAY)
    assert [(c[0], c[1], c[2], c[3]) for c in p.conflicts] == [("69000002", "8 Sep 2026", "P", "A")]
    sync(att, yes=True)
    assert csv_rows(att)[2][3] == "A" and ws_of(att).cell(row=3, column=4).value == "P"
    sync(att, answers=["x", "n"], monkeypatch=monkeypatch)
    assert csv_rows(att)[2][3] == "P"
    assert attendance.plan(att, TODAY).conflicts == []


def test_only_past_empty_sessions_are_offered(att, monkeypatch):
    p = attendance.plan(att, TODAY)
    assert p.empty_past == ["22 Sep 2026", "29 Sep 2026"]
    assert attendance.plan(att, date(2026, 9, 25)).empty_past == ["22 Sep 2026"]
    sync(att, yes=True)  # --yes never removes
    assert len(csv_rows(att)[0]) == 8
    sync(att, answers=["y"], monkeypatch=monkeypatch)
    assert csv_rows(att)[0] == ["Student ID", "Name"] + SESSIONS[:3] + ["total (3pts)"]
    assert header(ws_of(att)) == ["Student ID", "Name"] + SESSIONS[:3] + ["total (3pts)"]


def test_first_sync_proposes_the_newer_file(att, monkeypatch):
    attendance.baseline_path(att).unlink()
    edit_csv(att, lambda r: r[1].__setitem__(2, "A"))  # the csv is newer
    p = attendance.plan(att, TODAY)
    assert p.first_run and p.newer == "csv" and p.copies == [("69000001", "1 Sep 2026", "P", "A", "xlsx")]
    assert attendance.warnings(att)  # needs a question: not applied on open
    sync(att, answers=["n", "n"], monkeypatch=monkeypatch)  # empty sessions: n; apply: n
    assert ws_of(att).cell(row=2, column=3).value == "P"
    sync(att, answers=["n", "y"], monkeypatch=monkeypatch)
    assert ws_of(att).cell(row=2, column=3).value == "A"


def test_update_keeps_columns_and_follows_the_roster(att):
    delete_xlsx_col(att, "29 Sep 2026")
    sync(att, yes=True)
    roster = next((att / "course_info").glob("*_student_info.csv"))
    roster.write_text(roster.read_text(encoding="utf-8") + "69000004,นาย ง  สี่,G1\n", encoding="utf-8")
    grader._attendance_rows(att, xlsx_of(att), pd.read_csv(roster, dtype=str),
                            grader.calculate_class_dates("2026-06-15", "อ.", ""))
    rows = csv_rows(att)
    assert rows[0][-1] == "total (4pts)" and rows[-1][:2] == ["69000004", "นาย ง  สี่"] and rows[-1][-1] == "0"
    ws = ws_of(att)
    assert header(ws)[-1] == "total (4pts)" and ws.cell(row=5, column=1).value == "69000004"
    assert "C5:F5" in ws.cell(row=5, column=7).value
    assert attendance.plan(att, TODAY).has_changes() is False


def test_placeholder_sessions_get_dates_on_update(course):
    reg = pd.read_csv(next((course / "course_info").glob("*_student_info.csv")), dtype=str)
    path = course / "data" / "2026_S1_Test_SEC_51_attendance.xlsx"
    grader.save_attendance_excel(path, reg, ["a1", "a2", "a3"])
    attendance.make_csv(course)
    grader._attendance_rows(course, path, reg, ["1 Sep 2026", "8 Sep 2026", "15 Sep 2026", "22 Sep 2026"])
    assert csv_rows(course)[0][2:5] == ["1 Sep 2026", "8 Sep 2026", "15 Sep 2026"]
    assert header(ws_of(course))[2:5] == ["1 Sep 2026", "8 Sep 2026", "15 Sep 2026"]


def test_opening_twice_writes_nothing(att):
    tui_api.get_course_data(att)
    before = {p: p.stat().st_mtime_ns for p in (att / "data").iterdir()}
    n = len(list_entries(att))
    tui_api.get_course_data(att)
    tui_api.get_course_data(att)
    assert {p: p.stat().st_mtime_ns for p in (att / "data").iterdir()} == before
    assert len(list_entries(att)) == n


def test_lock_file_refuses_the_write(att):
    set_xlsx(att, "69000003", "15 Sep 2026", "A")
    lock = xlsx_of(att).parent / f"~${xlsx_of(att).name}"
    lock.write_text("x")
    before = xlsx_of(att).with_suffix(".csv").read_bytes()
    tui_api.get_course_data(att)
    sync(att, yes=True)
    assert xlsx_of(att).with_suffix(".csv").read_bytes() == before
    lock.unlink()
    sync(att, yes=True)
    assert csv_rows(att)[3][4] == "A"


def test_undo_restores_both_files(att):
    x0, c0 = xlsx_of(att).read_bytes(), xlsx_of(att).with_suffix(".csv").read_bytes()
    delete_xlsx_col(att, "29 Sep 2026")
    x1 = xlsx_of(att).read_bytes()
    sync(att, yes=True)
    assert xlsx_of(att).with_suffix(".csv").read_bytes() != c0
    grader.undo_course(course=str(att), yes=True)
    assert xlsx_of(att).read_bytes() == x1 and xlsx_of(att).with_suffix(".csv").read_bytes() == c0
    del x0


def test_undo_also_restores_the_baseline(att, monkeypatch):
    edit_csv(att, lambda rows: [r.pop(2) for r in rows])  # "1 Sep 2026", with marks
    sync(att, answers=["y", "n"], monkeypatch=monkeypatch)
    assert "1 Sep 2026" not in header(ws_of(att))
    grader.undo_course(course=str(att), yes=True)
    # back to before that sync: the csv lacks the column, the xlsx has it, and removing it is asked again
    p = attendance.plan(att, TODAY)
    assert [(r.session, r.deleted_in) for r in p.marked] == [("1 Sep 2026", "csv")]


def test_grade_tui_edit_keeps_both_files_and_the_total(att):
    tui_api.get_course_data(att)
    r = tui_api.update_student_score(att, "69000002", "29 Sep 2026", "P")
    assert r["status"] == "success", r
    assert csv_rows(att)[2][6] == "P" and csv_rows(att)[2][-1] == "2.8"
    assert ws_of(att).cell(row=3, column=7).value == "P"
    assert attendance.in_sync_fast(att) and attendance.warnings(att) == []
    r = tui_api.update_column_score(att, "22 Sep 2026", "X")
    assert r["status"] == "success", r
    assert [x[-1] for x in csv_rows(att)[1:]] == ["4", "3.8", "4"]
    assert attendance.in_sync_fast(att)
