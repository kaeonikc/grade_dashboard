"""`grader web-pull` (web/SPEC.md "Mac commands", REQUIREMENTS S11) and the optional `entry=` arguments."""
import json
import os
import re
import shutil

import openpyxl
import pytest
import yaml

from src import attendance, history, tui_api
from test_web_push import make_unsynced_pair, _set_xlsx
from webtools import (F1, JOURNALS, P1, csv_cell, grader_cli, grader_proc, journal_text, mirror, put_archive,
                      put_journal, tree_sha)

HW = f"data/{P1}_homework.csv"
ATT = f"data/{P1}_attendance.csv"
MID = f"data/{P1}_midterm.csv"
CFG = f"course_info/{P1}_config.yaml"

SAMPLE = """By column:
  20 Jan 2027   2 to apply
  hw2           1 to apply · 1 conflict · 1 skipped
  weights       1 to apply

To apply (Mac value → web value):
  9900000001 20 Jan 2027: (blank) → P
  9900000002 20 Jan 2027: (blank) → A
  9900000002 hw2: (blank) → 8
  weights: homework 0.2 → 0.25, final 0.4 → 0.35

Not applied (Mac values kept):
  conflict: 9900000001 hw2 — Mac 7, web 9 (not applied)
  skipped: 9900000099 hw2 — student not in Mac files

Already applied: 0 · No change: 0 · Replaced by a later edit: 1

Preview only: nothing was written. Apply with: grader web-pull --yes"""


def pushed(course, root):
    code, out = grader_cli(["web-push", str(course)])
    assert code == 0, out
    return mirror(root, course)


def mac_edit(course, sid, col, value, original):
    r = tui_api.apply_edits(course, [{"student_id": sid, "column": col, "value": value, "original": original}])
    assert r["status"] == "success" and r["saved"] == 1, r


def pull(course, *extra, cwd=None):
    return grader_cli(["web-pull", str(course), *extra], cwd=cwd)


def entries(course):
    return history.list_entries(course)


def no_bak(shas):
    return {k: v for k, v in shas.items() if not k.endswith(".bak")}


def archives(m):
    d = m / "archive"
    return sorted(p for p in d.glob("*.jsonl")) if d.is_dir() else []


def keys_of(path):
    out = []
    for line in path.read_text(encoding="utf-8").splitlines()[1:]:
        o = json.loads(line)
        out.append(f"{o['batch']}#{o['seq']}")
    return out


# ---- S11 -----------------------------------------------------------------------------------------------------------
def test_ac11_1_preview_matches_sample_session(web_course, drive_root):
    m = pushed(web_course, drive_root)
    mac_edit(web_course, "9900000001", "hw2", "7", "")
    put_journal(m, "seven_lines.jsonl")
    before, drive_before = tree_sha(web_course), tree_sha(drive_root)
    code, out = pull(web_course)
    assert code == 0, out
    assert "By column:" in out, out
    assert out[out.index("By column:"):].rstrip("\n") == SAMPLE
    assert "   Journal: 7 edits not pulled · last saved 20 Jan 2027 09:45" in out.splitlines()
    assert tree_sha(web_course) == before
    assert tree_sha(drive_root) == drive_before


def test_ac11_1_preview_same_with_and_without_tty(web_course, drive_root, tmp_path):
    m = pushed(web_course, drive_root)
    mac_edit(web_course, "9900000001", "hw2", "7", "")
    put_journal(m, "seven_lines.jsonl")
    before = tree_sha(web_course)
    plain = grader_proc(["web-pull", web_course], cwd=tmp_path)
    tty = grader_proc(["web-pull", web_course], cwd=tmp_path, tty=True)
    assert plain.returncode == 0 and tty.returncode == 0, plain.stdout + plain.stderr
    assert plain.stdout == tty.stdout
    assert plain.stdout[plain.stdout.index("By column:"):].rstrip("\n") == SAMPLE
    assert tree_sha(web_course) == before


def test_ac11_2_yes_applies_with_totals_and_one_history_entry(web_course, drive_root):
    m = pushed(web_course, drive_root)
    put_journal(m, "four_lines.jsonl")
    n0 = len(entries(web_course))
    code, out = pull(web_course, "--yes")
    assert code == 0, out
    c = web_course
    assert csv_cell(c / HW, "9900000001", "hw2") == "9"
    assert csv_cell(c / HW, "9900000001", "total (30pts)") == "17"
    assert csv_cell(c / MID, "9900000002", "midterm_exam") == "20"
    assert csv_cell(c / MID, "9900000002", "total (30pts)") == "20"
    assert csv_cell(c / ATT, "9900000001", "20 Jan 2027") == "P"
    assert csv_cell(c / ATT, "9900000002", "20 Jan 2027") == "A"
    new = entries(c)[n0:]
    assert len(new) == 1, new
    assert (new[0]["tool"], new[0]["action"]) == ("grader", "web-pull")
    assert set(new[0]["files"]) >= {HW, MID, ATT}
    code, hist = grader_cli(["history", "--course", str(c)])
    assert sum("web-pull" in l for l in hist.splitlines()) == 1
    assert f"Undo with: grader undo {new[0]['id']}" in out


def test_ac11_3_one_undo_restores_every_file(web_course_att, drive_root):
    c = web_course_att
    m = pushed(c, drive_root)
    put_journal(m, "seven_lines.jsonl")
    watched = lambda: no_bak({k: v for k, v in tree_sha(c).items()
                              if k.startswith(("data/", "course_info/")) or k == "history/attendance_sync.json"})
    before = watched()
    code, out = pull(c, "--yes")
    assert code == 0, out
    after = watched()
    changed = {k for k in before if before[k] != after.get(k)}
    assert {HW, ATT, f"data/{P1}_attendance.xlsx", CFG} <= changed, changed
    code, out = grader_cli(["undo", "--course", str(c), "--yes"])
    assert code == 0, out
    now = watched()
    assert {k: now.get(k) for k in before} == before


def test_ac11_4_mac_change_is_conflict(web_course, drive_root):
    m = pushed(web_course, drive_root)
    mac_edit(web_course, "9900000001", "hw2", "7", "")
    put_journal(m, "seven_lines.jsonl")
    code, out = pull(web_course, "--yes")
    assert code == 0, out
    assert "  conflict: 9900000001 hw2 — Mac 7, web 9 (not applied)" in out.splitlines()
    assert csv_cell(web_course / HW, "9900000001", "hw2") == "7"
    assert csv_cell(web_course / HW, "9900000002", "hw2") == "8"


def test_ac11_5_unknown_student_or_column_skipped(web_course, drive_root):
    m = pushed(web_course, drive_root)
    put_journal(m, "unknown.jsonl")
    code, out = pull(web_course, "--yes")
    assert code == 0, out
    lines = out.splitlines()
    assert "  skipped: 9900000099 hw2 — student not in Mac files" in lines
    assert "  skipped: 9900000001 hw9 — column not in Mac files" in lines
    assert "  skipped: 9900000001 total (30pts) — total column" in lines
    assert csv_cell(web_course / HW, "9900000002", "hw3") == "6"
    assert csv_cell(web_course / HW, "9900000001", "total (30pts)") == "8"


def test_ac11_6_net_change_against_push_value(web_course, drive_root):
    m = pushed(web_course, drive_root)
    put_journal(m, "net_change.jsonl")
    code, out = pull(web_course)
    assert code == 0, out
    assert "  9900000001 hw2: (blank) → 8" in out.splitlines()
    assert "9900000001 hw2: (blank) → 9" not in out
    assert "Replaced by a later edit: 2" in out
    code, out = pull(web_course, "--yes")
    assert code == 0, out
    assert csv_cell(web_course / HW, "9900000001", "hw2") == "8"


def test_ac11_7_equal_value_is_already_applied(web_course, drive_root):
    m = pushed(web_course, drive_root)
    put_journal(m, "duplicates.jsonl")
    code, out = pull(web_course, "--yes")  # duplicate lines from a retried save: applied once, no conflict
    assert code == 0, out
    assert "conflict" not in out
    assert csv_cell(web_course / HW, "9900000001", "hw2") == "9"
    assert csv_cell(web_course / HW, "9900000001", "total (30pts)") == "17"
    # edge 23: the Mac already holds the value the web set
    m2 = m
    put_journal(m2, "four_lines.jsonl")
    mac_edit(web_course, "9900000002", "midterm_exam", "20", "18")
    code, out = pull(web_course)
    assert code == 0, out
    assert "already applied: 9900000002 midterm_exam = 20" in out
    assert "conflict" not in out


def test_ac11_8_xlsx_mirror_keeps_formulas_validation_and_colours(web_course_att, drive_root):
    c = web_course_att
    x = c / "data" / f"{P1}_attendance.xlsx"

    def shape():
        ws = openpyxl.load_workbook(x).worksheets[0]
        dv = sorted((d.formula1, str(d.sqref)) for d in ws.data_validations.dataValidation)
        cf = sorted((str(r.sqref), tuple(tuple(rule.formula or ()) for rule in r.rules)) for r in ws.conditional_formatting)
        return ws, dv, cf
    _, dv0, cf0 = shape()
    m = pushed(c, drive_root)
    put_journal(m, "four_lines.jsonl")
    code, out = pull(c, "--yes")
    assert code == 0, out
    ws, dv1, cf1 = shape()
    head = [cell.value for cell in ws[1]]
    col = head.index("20 Jan 2027") + 1
    tot = next(i for i, h in enumerate(head, 1) if str(h).lower().startswith("total"))
    vals = {str(ws.cell(row=r, column=1).value): ws.cell(row=r, column=col).value for r in range(2, ws.max_row + 1)}
    assert vals["9900000001"] == "P" and vals["9900000002"] == "A"
    for r in range(2, ws.max_row + 1):
        assert str(ws.cell(row=r, column=tot).value).startswith("="), r
    assert dv1 == dv0 and cf1 == cf0
    assert csv_cell(c / ATT, "9900000001", "20 Jan 2027") == "P"


def test_ac11_9_archive_on_drive_and_in_history(web_course, drive_root):
    m = pushed(web_course, drive_root)
    j = put_journal(m, "four_lines.jsonl")
    jsha = tree_sha(m)["web_edits.jsonl"]
    n0 = len(entries(web_course))
    code, out = pull(web_course, "--yes")
    assert code == 0, out
    arch = archives(m)
    assert len(arch) == 1 and re.fullmatch(r"web_edits_\d{8}-\d{6}\.jsonl", arch[0].name), arch
    assert f"📦 Archived on Drive: archive/{arch[0].name} (4 lines)" in out
    lines = [json.loads(l) for l in arch[0].read_text(encoding="utf-8").splitlines()]
    assert lines[0]["kind"] == "pull" and len(lines) == 5
    assert all(l["result"] == "applied" for l in lines[1:])
    e = entries(web_course)[n0:]
    assert len(e) == 1
    copy = web_course / "history" / e[0]["id"] / arch[0].name
    assert copy.read_bytes() == arch[0].read_bytes()
    assert not any(f.endswith(".jsonl") for f in e[0]["files"])
    assert tree_sha(m)["web_edits.jsonl"] == jsha  # the Mac never rewrites the journal
    code, out = pull(web_course)
    assert "✅ Nothing to pull: no web edits since the last pull." in out
    with open(j, "a", encoding="utf-8") as f:  # a line saved on the web after that read
        f.write(json.dumps({"batch": "b-20270120-160000-ab12", "seq": 0, "ts": "2027-01-20T16:00:00+07:00",
                            "kind": "cell", "push": "p-20270119-201500", "file": HW, "student_id": "9900000001",
                            "column": "hw3", "value": "6", "original": ""}, ensure_ascii=False) + "\n")
    code, out = pull(web_course)
    assert code == 0, out
    assert "Journal: 1 edit not pulled · last saved 20 Jan 2027 16:00" in out
    assert "  9900000001 hw3: (blank) → 6" in out.splitlines()


def test_ac11_10_second_pull_nothing(web_course, drive_root):
    m = pushed(web_course, drive_root)
    put_journal(m, "four_lines.jsonl")
    assert pull(web_course, "--yes")[0] == 0
    n = len(entries(web_course))
    code, out = pull(web_course, "--yes")
    assert code == 0
    assert "✅ Nothing to pull: no web edits since the last pull." in out
    assert len(entries(web_course)) == n


def _cfg(course):
    return yaml.safe_load((course / CFG).read_text(encoding="utf-8"))


def test_ac11_11_config_applied_when_mac_unchanged(web_course, drive_root):
    m = pushed(web_course, drive_root)
    put_journal(m, "config_weights.jsonl")
    before = _cfg(web_course)
    code, out = pull(web_course, "--yes")
    assert code == 0, out
    after = _cfg(web_course)
    assert after["weights"] == {"attendance": 0.1, "homework": 0.25, "midterm": 0.3, "final": 0.35}
    for k in before:
        if k != "weights":
            assert after[k] == before[k], k


def test_ac11_12_config_conflict_when_mac_changed(web_course, drive_root):
    m = pushed(web_course, drive_root)
    r = tui_api.update_course_config(web_course, weights_dict={"attendance": 0.1, "homework": 0.15, "midterm": 0.3,
                                                               "final": 0.4})
    assert r["status"] == "success"
    put_journal(m, "config_weights.jsonl")
    code, out = pull(web_course, "--yes")
    assert code == 0, out
    line = next((l for l in out.splitlines() if l.startswith("  conflict: weights — Mac ")), None)
    assert line, out
    mac, web = line.split(", web ", 1) if ", web " in line else line.split(" web ", 1)
    assert "homework 0.15" in mac and "homework 0.25" in web
    assert line.endswith("; apply by hand if wanted")
    assert _cfg(web_course)["weights"]["homework"] == 0.15
    # the same rule for grade boundaries
    r = tui_api.update_course_config(web_course, boundaries_dict={"A": 82, "B+": 75, "B": 70, "C+": 65, "C": 60,
                                                                  "D+": 55, "D": 50})
    put_journal(m, "config_boundaries.jsonl")
    code, out = pull(web_course, "--yes")
    assert code == 0, out
    line = next((l for l in out.splitlines() if l.startswith("  conflict: grade boundaries — Mac ")), None)
    assert line and "A 82" in line and "A 78" in line and line.endswith("; apply by hand if wanted"), out
    assert _cfg(web_course)["grade_boundaries"]["A"] == 82


def test_ac11_13_invalid_line_applies_nothing(web_course, drive_root):
    m = pushed(web_course, drive_root)
    put_journal(m, "truncated.jsonl")
    before, dbefore = tree_sha(web_course), tree_sha(drive_root)
    code, out = pull(web_course, "--yes")
    assert code == 1
    assert "❌ web_edits.jsonl has lines that aren't valid JSON: 3. Nothing was applied, archived or pushed." in out
    assert tree_sha(web_course) == before and tree_sha(drive_root) == dbefore


def test_ac11_14_other_course_refused(web_course, drive_root):
    m = pushed(web_course, drive_root)
    put_journal(m, "four_lines.jsonl")
    cfg = web_course / CFG
    cfg.write_text(cfg.read_text(encoding="utf-8").replace("course_id: MOCK", "course_id: PHYS1120")
                   .replace("term: 2027_S1", "term: 2027_S2"), encoding="utf-8")
    before = tree_sha(web_course)
    code, out = pull(web_course, "--yes")
    assert code == 1
    assert ("❌ The Drive copy is another course: Drive has MOCK 2027_S1, this folder is PHYS1120 2027_S2. "
            "Nothing was pulled.") in out
    assert tree_sha(web_course) == before


def test_ac11_15_locked_xlsx_or_unsynced_pair_refuses_yes(web_course_att, drive_root):
    c = web_course_att
    m = pushed(c, drive_root)
    put_journal(m, "four_lines.jsonl")
    lock = c / "data" / f"~${P1}_attendance.xlsx"
    lock.write_bytes(b"lock")
    before, dbefore = tree_sha(c), tree_sha(drive_root)
    code, out = pull(c, "--yes")
    assert code == 1
    assert "To apply (Mac value → web value):" in out
    assert f"❌ Not applied: data/{P1}_attendance.xlsx is open in Excel." in out
    assert tree_sha(c) == before and tree_sha(drive_root) == dbefore
    assert pull(c)[0] == 0  # the preview still works
    lock.unlink()
    make_unsynced_pair(c)
    before = tree_sha(c)
    code, out = pull(c, "--yes")
    assert code == 1
    assert "❌ Not applied: the attendance xlsx and csv differ; run `grader sync` first" in out
    assert tree_sha(c) == before and tree_sha(drive_root) == dbefore


def test_edge30_lock_ignored_without_attendance_lines(web_course_att, drive_root):
    c = web_course_att
    m = pushed(c, drive_root)
    put_journal(m, "net_change.jsonl")
    (c / "data" / f"~${P1}_attendance.xlsx").write_bytes(b"lock")
    code, out = pull(c, "--yes")
    assert code == 0, out
    assert csv_cell(c / HW, "9900000001", "hw2") == "8"


def test_ac11_16_cwd_default(web_course, drive_root):
    m = pushed(web_course, drive_root)
    put_journal(m, "four_lines.jsonl")
    code, out = grader_cli(["web-pull"], cwd=web_course)
    assert code == 0, out
    assert "To apply (Mac value → web value):" in out


def test_ac11_16_usage_error_outside_a_course(tmp_path):
    d = tmp_path / "elsewhere"
    d.mkdir()
    r = grader_proc(["web-pull"], cwd=d)
    assert r.returncode == 2
    assert "usage: grader web-pull [-h] [-y] [--drive-root DIR] [--debug] [COURSE]" in r.stderr


# ---- edges ---------------------------------------------------------------------------------------------------------
def _five_lines():
    four = journal_text("four_lines.jsonl")
    fifth = json.dumps({"batch": "b-20270120-170000-ef01", "seq": 0, "ts": "2027-01-20T17:00:00+07:00",
                        "kind": "cell", "push": "p-20270119-201500", "file": HW, "student_id": "9900000003",
                        "column": "hw2", "value": "6", "original": ""}, ensure_ascii=False, separators=(",", ":"))
    return (four + fifth + "\n").splitlines(keepends=True)


def test_edge15_stale_local_journal_never_loses_cloud_lines(web_course, drive_root):
    m = pushed(web_course, drive_root)
    lines = _five_lines()
    (m / "web_edits.jsonl").write_text("".join(lines[:3]), encoding="utf-8")  # the local copy lags the cloud
    code, out = pull(web_course, "--yes")
    assert code == 0, out
    (m / "web_edits.jsonl").write_text("".join(lines), encoding="utf-8")  # Drive for Desktop catches up
    code, out = pull(web_course)
    assert "Journal: 2 edits not pulled" in out, out
    code, out = pull(web_course, "--yes")
    assert code == 0, out
    assert csv_cell(web_course / MID, "9900000002", "midterm_exam") == "20"
    assert csv_cell(web_course / HW, "9900000003", "hw2") == "6"
    all_keys = [k for a in archives(m) for k in keys_of(a)]
    want = [f"{json.loads(l)['batch']}#{json.loads(l)['seq']}" for l in lines]
    assert sorted(all_keys) == sorted(want)  # every key archived exactly once


def test_edge15_push_during_unsynced_save_keeps_lines(web_course, drive_root):
    m = pushed(web_course, drive_root)
    assert pull(web_course)[1].count("Nothing to pull") == 1
    old = json.loads((m / "push.json").read_text(encoding="utf-8"))["push_id"]
    assert grader_cli(["web-push", str(web_course)])[0] == 0  # pushed while a web save hadn't synced yet
    two = "".join(journal_text("four_lines.jsonl", push=old).splitlines(keepends=True)[2:])
    (m / "web_edits.jsonl").write_text(two, encoding="utf-8")  # that save arrives with the old push id
    code, out = grader_cli(["web-push", str(web_course)])
    assert code == 1
    assert "2 web edits haven't been pulled yet" in out
    code, out = pull(web_course, "--yes")
    assert code == 0, out
    assert csv_cell(web_course / HW, "9900000001", "hw2") == "9"
    assert csv_cell(web_course / MID, "9900000002", "midterm_exam") == "20"


def test_edge15_mac_never_writes_the_journal(web_course, drive_root):
    m = pushed(web_course, drive_root)
    put_journal(m, "three_unpulled.jsonl")
    sha = lambda: tree_sha(m)["web_edits.jsonl"]
    s0 = sha()
    assert grader_cli(["web-push", str(web_course)])[0] == 1
    assert sha() == s0
    assert grader_cli(["web-push", str(web_course), "--force"])[0] == 0
    assert sha() == s0
    with open(m / "web_edits.jsonl", "a", encoding="utf-8") as f:
        f.write(journal_text("four_lines.jsonl"))
    s1 = sha()
    assert pull(web_course)[0] == 0
    assert sha() == s1
    assert pull(web_course, "--yes")[0] == 0
    assert sha() == s1


def test_edge15_prints_count_and_last_save(web_course, drive_root):
    m = pushed(web_course, drive_root)
    put_journal(m, "seven_lines.jsonl")
    code, out = pull(web_course)
    assert "   Journal: 7 edits not pulled · last saved 20 Jan 2027 09:45" in out.splitlines()
    assert "(grade_web shows the same two numbers. If not, wait for Drive to sync and run again.)" in out


def test_edge17_conflict_copies_refuse_pull(web_course, drive_root):
    m = pushed(web_course, drive_root)
    put_journal(m, "four_lines.jsonl")
    shutil.copyfile(m / HW, m / "data" / f"{P1}_homework (1).csv")
    before = tree_sha(web_course)
    code, out = pull(web_course, "--yes")
    assert code == 1
    assert f"❌ Drive has duplicate files in grade_web/{F1}/" in out and f"{P1}_homework (1).csv" in out
    assert tree_sha(web_course) == before


@pytest.mark.skipif(os.geteuid() == 0, reason="root can read a chmod 000 file")
def test_edge19_unreadable_journal_refuses_pull(web_course, drive_root):
    m = pushed(web_course, drive_root)
    j = put_journal(m, "four_lines.jsonl")
    j.chmod(0)
    try:
        code, out = pull(web_course, "--yes")
    finally:
        j.chmod(0o644)
    assert code == 1
    assert f"❌ Couldn't read grade_web/{F1}/web_edits.jsonl (" in out
    assert "Nothing to pull" not in out


def test_edge24_edited_back_is_no_change(web_course, drive_root):
    m = pushed(web_course, drive_root)
    put_journal(m, "net_change.jsonl")
    hw_before = (web_course / HW).read_bytes()
    code, out = pull(web_course, "--yes")
    assert code == 0, out
    assert "  no change: 9900000002 hw3 (edited back to 7)" in out.splitlines()
    assert csv_cell(web_course / HW, "9900000002", "hw3") == "7"
    assert hw_before != (web_course / HW).read_bytes()  # hw2 of 9900000001 did change


def test_edge26_header_points_changed_still_applies(web_course, drive_root):
    m = pushed(web_course, drive_root)
    hw = web_course / HW
    hw.write_bytes(hw.read_bytes().replace(b"hw2 (10pts)", b"hw2 (15pts)"))
    put_journal(m, "four_lines.jsonl")
    code, out = pull(web_course, "--yes")
    assert code == 0, out
    assert csv_cell(hw, "9900000001", "hw2 (15pts)") == "9"
    assert "skipped" not in out


def test_edge27_new_mac_student_untouched(web_course, drive_root):
    m = pushed(web_course, drive_root)
    for rel, row in ((f"course_info/{P1}_student_info.csv", "9900000004,นายใหม่ เพิ่มมา,SEC2\n"),
                     (HW, "9900000004,นายใหม่ เพิ่มมา,6,,,6\n")):
        p = web_course / rel
        p.write_text(p.read_text(encoding="utf-8") + row, encoding="utf-8")
    put_journal(m, "four_lines.jsonl")
    code, out = pull(web_course, "--yes")
    assert code == 0, out
    text = (web_course / HW).read_text(encoding="utf-8")
    assert "9900000004,นายใหม่ เพิ่มมา,6,,,6" in text.splitlines()


@pytest.mark.skipif(os.geteuid() == 0, reason="root ignores directory permissions")
def test_edge28_crash_before_archive_rerun_only_archives(web_course, drive_root):
    m = pushed(web_course, drive_root)
    put_journal(m, "four_lines.jsonl")
    (m / "archive").mkdir()
    (m / "archive").chmod(0o555)
    try:
        code, out = pull(web_course, "--yes")
    finally:
        (m / "archive").chmod(0o755)
    assert code == 1
    assert "⚠️  Applied on the Mac (undo with: grader undo " in out
    assert "but not archived on Drive" in out
    assert "Run grader web-pull --yes again" in out
    assert csv_cell(web_course / HW, "9900000001", "hw2") == "9"
    n = len(entries(web_course))
    code, out = pull(web_course, "--yes")
    assert code == 0, out
    assert "✅ Nothing to change on the Mac (4 already applied)." in out
    assert len(archives(m)) == 1
    assert len(entries(web_course)) == n


def test_edge32_several_config_lines_net_change(web_course, drive_root):
    m = pushed(web_course, drive_root)
    put_journal(m, "config_weights_net.jsonl")
    code, out = pull(web_course)
    assert "  weights: homework 0.2 → 0.25, final 0.4 → 0.35" in out.splitlines(), out
    assert "Replaced by a later edit: 1" in out
    code, out = pull(web_course, "--yes")
    assert code == 0, out
    assert _cfg(web_course)["weights"] == {"attendance": 0.1, "homework": 0.25, "midterm": 0.3, "final": 0.35}


def test_from_web_journal_pulls_cleanly(web_course, drive_root):
    m = pushed(web_course, drive_root)
    put_journal(m, "from_web.jsonl")
    code, out = pull(web_course, "--yes")
    assert code == 0, out
    assert "conflict" not in out and "skipped:" not in out
    assert csv_cell(web_course / HW, "9900000001", "hw2") == "9"
    assert csv_cell(web_course / HW, "9900000002", "hw2") == "8"
    assert csv_cell(web_course / MID, "9900000002", "midterm_exam") == "20"


def test_journal_contract_matches_expected_json(web_course, drive_root):
    """Python reads every fixture journal as expected.json says (node checks the same file)."""
    exp = json.loads((JOURNALS / "expected.json").read_text(encoding="utf-8"))
    m = pushed(web_course, drive_root)
    put_archive(m, exp["archive"])
    months = {"01": "Jan"}
    for name, e in exp["journals"].items():
        put_journal(m, name)
        code, out = pull(web_course)
        if e["bad"]:
            assert code == 1, name
            assert f"aren't valid JSON: {', '.join(map(str, e['bad']))}." in out, (name, out)
        elif e["unpulled"] == 0:
            assert "Nothing to pull" in out, name
        else:
            assert code == 0, (name, out)
            ts = e["last_saved"]
            when = f"{int(ts[8:10])} {months[ts[5:7]]} {ts[:4]} {ts[11:16]}"
            word = "edit" if e["unpulled"] == 1 else "edits"
            assert f"Journal: {e['unpulled']} {word} not pulled · last saved {when}" in out, (name, out)


# ---- the optional entry= arguments (ADR 002) ---------------------------------------------------------------------
def test_apply_edits_entry_captures_into_one_entry(web_course_att):
    c = web_course_att
    n0 = len(entries(c))
    e = history.Entry(c, "grader", "web-pull")
    r = tui_api.apply_edits(c, [
        {"student_id": "9900000001", "column": "hw2", "value": "9", "original": ""},
        {"student_id": "9900000002", "column": "20 Jan 2027", "value": "A", "original": ""},
    ], entry=e)
    assert r["status"] == "success" and r["saved"] == 2, r
    assert len(entries(c)) == n0  # the caller closes the entry
    e.close(cells=2)
    new = entries(c)[n0:]
    assert len(new) == 1
    assert set(new[0]["files"]) >= {HW, ATT, f"data/{P1}_attendance.xlsx", "history/attendance_sync.json"}
    assert csv_cell(c / HW, "9900000001", "total (30pts)") == "17"
    assert attendance.in_sync_fast(c)


def test_auto_sync_entry_param(web_course_att):
    c = web_course_att
    _set_xlsx(c, "9900000001", "20 Jan 2027", "P")  # one-sided change: auto-sync copies it, no question
    n0 = len(entries(c))
    e = history.Entry(c, "grader", "web-pull")
    warns = attendance.auto_sync(c, tool="grade-tui", entry=e)
    assert warns == []
    assert csv_cell(c / ATT, "9900000001", "20 Jan 2027") == "P"
    assert len(entries(c)) == n0
    e.close()
    new = entries(c)[n0:]
    assert len(new) == 1 and ATT in new[0]["files"] and "history/attendance_sync.json" in new[0]["files"]
