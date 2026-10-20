"""`grader web-push` (web/SPEC.md "Mac commands", REQUIREMENTS S1). Every test uses a tmp Drive root (conftest)."""
import json
import os
import re
import shutil

import openpyxl
import pytest

from conftest import build_attendance_xlsx
from src import history, tui_api
from webtools import (F1, F2, FIXTURES, P1, csv_cell, grader_cli, grader_proc, mirror, put_journal, tree_sha)

PUSHED = sorted(["course.json", "push.json", "web_edits.jsonl", f"course_info/{P1}_student_info.csv",
                 f"data/{P1}_attendance.csv", f"data/{P1}_final.csv", f"data/{P1}_homework.csv",
                 f"data/{P1}_midterm.csv"])


def push(course, *extra, cwd=None):
    return grader_cli(["web-push", str(course), *extra], cwd=cwd)


def test_ac1_1_push_copies_pushed_set(web_course_att, drive_root):
    c = web_course_att
    (c / "reports").mkdir()
    (c / "reports" / f"{P1}_final_grades.csv").write_text("Student ID\n9900000001\n", encoding="utf-8")
    history.snapshot(c, [c / "data" / f"{P1}_final.csv"], "grader", "test")
    code, out = push(c)
    assert code == 0, out
    m = mirror(drive_root, c)
    assert sorted(tree_sha(m)) == PUSHED
    for rel in PUSHED:
        if rel.startswith(("data/", "course_info/")):
            assert (m / rel).read_bytes() == (c / rel).read_bytes(), rel
    assert (m / "web_edits.jsonl").read_bytes() == b""
    rec = json.loads((m / "push.json").read_text(encoding="utf-8"))
    assert rec["course_id"] == "MOCK" and rec["course_name"] == "ฟิสิกส์จำลอง" and rec["term"] == "2027_S1"
    assert rec["mac_folder"] == F1
    assert re.fullmatch(r"\d{4}-\d\d-\d\dT\d\d:\d\d:\d\d[+-]\d\d:\d\d", rec["pushed_at"]), rec["pushed_at"]
    assert re.fullmatch(r"p-\d{8}-\d{6}", rec["push_id"])
    assert not list(m.rglob("*.xlsx")) and not (m / "history").exists() and not (m / "reports").exists()
    assert "✅ Pushed" in out


def test_ac1_2_push_never_writes_mac_folder(web_course_att, drive_root):
    c = web_course_att
    history.snapshot(c, [c / "data" / f"{P1}_midterm.csv"], "grader", "test")
    hw = c / "data" / f"{P1}_homework.csv"
    hw.write_bytes(hw.read_bytes().replace("ทดสอบ,8,,,8".encode(), "ทดสอบ,8,,,99".encode()))  # stale total
    before = tree_sha(c)
    code, out = push(c)
    assert code == 0, out
    assert tree_sha(c) == before
    # and after refusals too
    put_journal(mirror(drive_root, c), "three_unpulled.jsonl")
    code, _ = push(c)
    assert code == 1
    assert tree_sha(c) == before
    (c / "data" / f"{P1}_quiz.csv").write_bytes(b"Student ID,Name,q1\n9900000001,\xa1\xa2,1\n")
    before = tree_sha(c)
    code, _ = push(c, "--force")
    assert code == 1
    assert tree_sha(c) == before


def test_ac1_3_unpulled_lines_refuse_push(web_course, drive_root):
    assert push(web_course)[0] == 0
    m = mirror(drive_root, web_course)
    put_journal(m, "three_unpulled.jsonl")
    before = tree_sha(drive_root)
    code, out = push(web_course)
    assert code == 1
    assert "❌ 3 web edits haven't been pulled yet (last saved 20 Jan 2027 09:15):" in out
    tokens = [l.split() for l in out.splitlines()]
    for want in (["9900000001", "20", "Jan", "2027", "(blank)", "→", "P", "20", "Jan", "2027", "09:15"],
                 ["9900000002", "20", "Jan", "2027", "(blank)", "→", "A", "20", "Jan", "2027", "09:15"],
                 ["9900000002", "hw2", "(blank)", "→", "8.5", "20", "Jan", "2027", "09:15"]):
        assert want in tokens, (want, out)
    assert "grader web-push --force" in out
    assert "Nothing on Drive was changed." in out
    assert tree_sha(drive_root) == before


def test_ac1_4_force_archives_discarded_lines(web_course, drive_root):
    assert push(web_course)[0] == 0
    m = mirror(drive_root, web_course)
    j = put_journal(m, "three_unpulled.jsonl")
    jbytes = j.read_bytes()
    code, out = push(web_course, "--force")
    assert code == 0, out
    arch = list((m / "archive").glob("*"))
    assert len(arch) == 1
    assert re.fullmatch(r"web_edits_\d{8}-\d{6}_discarded\.jsonl", arch[0].name), arch[0].name
    assert f"archive/{arch[0].name}" in out
    lines = [json.loads(l) for l in arch[0].read_text(encoding="utf-8").splitlines()]
    assert lines[0]["kind"] == "discarded"
    assert len(lines) == 4 and all(l["result"] == "discarded" for l in lines[1:])
    assert {(l["student_id"], l["column"]) for l in lines[1:]} == {
        ("9900000001", "20 Jan 2027"), ("9900000002", "20 Jan 2027"), ("9900000002", "hw2")}
    assert j.read_bytes() == jbytes  # the journal is never rewritten by the Mac
    assert "✅ Pushed" in out
    code, out = push(web_course)  # unpulled is now 0: a plain push goes through
    assert code == 0, out
    assert "0 edits waiting to be pulled" in out


def _set_xlsx(course, sid, session, value):
    x = next((course / "data").glob("*_attendance.xlsx"))
    wb = openpyxl.load_workbook(x)
    ws = wb.worksheets[0]
    head = [c.value for c in ws[1]]
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(row=r, column=1).value) == sid:
            ws.cell(row=r, column=head.index(session) + 1).value = value
    wb.save(x)
    t = os.path.getmtime(x) + 5
    os.utime(x, (t, t))


def _set_csv(course, sid, session, value):
    p = next((course / "data").glob("*_attendance.csv"))
    lines = p.read_text(encoding="utf-8").splitlines()
    head = lines[0].split(",")
    i = head.index(session)
    for n, line in enumerate(lines):
        cells = line.split(",")
        if cells[0] == sid:
            cells[i] = value
            lines[n] = ",".join(cells)
    p.write_text("\n".join(lines) + "\n", encoding="utf-8")
    t = os.path.getmtime(p) + 10
    os.utime(p, (t, t))


def make_unsynced_pair(course):
    """9900000002 / 20 Jan 2027: P in the xlsx, A in the csv, both changed since the baseline (needs grader sync)."""
    _set_xlsx(course, "9900000002", "20 Jan 2027", "P")
    _set_csv(course, "9900000002", "20 Jan 2027", "A")


def test_ac1_5_attendance_needing_sync_refuses(web_course_att, drive_root):
    make_unsynced_pair(web_course_att)
    before = tree_sha(web_course_att)
    code, out = push(web_course_att)
    assert code == 1
    assert "run `grader sync` first" in out
    assert any("9900000002" in l and "20 Jan 2027" in l and "xlsx P" in l and "csv A" in l for l in out.splitlines()), out
    assert tree_sha(drive_root) == {}
    assert not (drive_root / "grade_courses").exists() and not (drive_root / "grade_web").exists()
    assert tree_sha(web_course_att) == before


def test_ac1_6_non_utf8_csv_refused(web_course, drive_root):
    hw = web_course / "data" / f"{P1}_homework.csv"
    hw.write_bytes(hw.read_bytes().decode("utf-8").encode("cp874"))
    code, out = push(web_course)
    assert code == 1
    assert f"❌ data/{P1}_homework.csv is not UTF-8 text" in out
    assert not (drive_root / "grade_courses").exists() and not (drive_root / "grade_web").exists()


def test_ac1_7_stale_drive_file_removed(web_course, drive_root):
    quiz = web_course / "data" / f"{P1}_quiz.csv"
    quiz.write_text("Student ID,Name,q1\n9900000001,นายสมมติ ทดสอบ,3\n", encoding="utf-8")
    assert push(web_course)[0] == 0
    m = mirror(drive_root, web_course)
    assert (m / "data" / quiz.name).exists()
    (m / "archive").mkdir()
    (m / "archive" / "web_edits_20270101-000000.jsonl").write_text('{"kind":"pull"}\n', encoding="utf-8")
    keep = tree_sha(m / "archive")
    quiz.unlink()
    code, out = push(web_course)
    assert code == 0, out
    assert f"Removed from Drive: data/{P1}_quiz.csv" in out
    assert sorted(k for k in tree_sha(m) if not k.startswith("archive/")) == PUSHED
    assert tree_sha(m / "archive") == keep


def test_ac1_8_env_root_isolates_tests(web_course, drive_root, tmp_path, monkeypatch):
    home = tmp_path / "home"
    cloud = home / "Library" / "CloudStorage" / "GoogleDrive-fake.teacher@g.cmru.ac.th" / "My Drive"
    cloud.mkdir(parents=True)
    monkeypatch.setenv("HOME", str(home))
    code, out = push(web_course)
    assert code == 0, out
    assert list(cloud.iterdir()) == []
    assert mirror(drive_root, web_course).is_dir()


def test_ac1_8_default_root_resolution(web_course, tmp_path):
    home = tmp_path / "home"
    cs = home / "Library" / "CloudStorage"
    cs.mkdir(parents=True)
    # zero matches
    r = grader_proc(["web-push", web_course], cwd=tmp_path, home=home, unset=["GRADE_WEB_DRIVE_ROOT"])
    assert r.returncode == 1, r.stdout + r.stderr
    assert "❌ No Google Drive folder for an @g.cmru.ac.th account under ~/Library/CloudStorage/." in r.stdout
    # one match: used
    one = cs / "GoogleDrive-fake.teacher@g.cmru.ac.th" / "My Drive"
    one.mkdir(parents=True)
    r = grader_proc(["web-push", web_course], cwd=tmp_path, home=home, unset=["GRADE_WEB_DRIVE_ROOT"])
    assert r.returncode == 0, r.stdout + r.stderr
    assert (one / "grade_courses" / "mac" / F1 / "course.json").exists()
    # two matches: refused
    (cs / "GoogleDrive-other.teacher@g.cmru.ac.th" / "My Drive").mkdir(parents=True)
    r = grader_proc(["web-push", web_course], cwd=tmp_path, home=home, unset=["GRADE_WEB_DRIVE_ROOT"])
    assert r.returncode == 1
    assert "❌ Found 2 Google Drive folders:" in r.stdout and "--drive-root" in r.stdout


def test_ac1_9_missing_drive_root(web_course, tmp_path):
    gone = tmp_path / "no-such-drive"
    code, out = push(web_course, "--drive-root", gone)
    assert code == 1
    assert f"❌ Drive folder not found: {gone}" in out
    assert "Nothing was written." in out
    assert not gone.exists()


def test_ac1_10_xlsx_score_file_refused(web_course_att, drive_root):
    q = web_course_att / "data" / f"{P1}_quiz.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["Student ID", "Name", "q1 (5pts)"])
    wb.active.append(["9900000001", "นายสมมติ ทดสอบ", 3])
    wb.save(q)
    code, out = push(web_course_att)
    assert code == 1
    assert f"❌ grade_web reads score files as CSV only: data/{P1}_quiz.xlsx" in out
    assert "attendance.xlsx" not in out.split("CSV only:")[1].splitlines()[0]
    assert not (drive_root / "grade_courses").exists() and not (drive_root / "grade_web").exists()


def test_ac1_11_cwd_default(web_course, drive_root):
    code, out = grader_cli(["web-push"], cwd=web_course)
    assert code == 0, out
    assert (mirror(drive_root, web_course) / "course.json").exists()


def test_ac1_11_usage_error_outside_a_course(tmp_path):
    empty = tmp_path / "elsewhere"
    empty.mkdir()
    r = grader_proc(["web-push"], cwd=empty)
    assert r.returncode == 2
    assert "usage: grader web-push [-h] [--force] [--semester NAME] [--drive-root DIR]" in r.stderr
    assert "grader web-push: error: no course folder here" in r.stderr


def test_ac3_1_course_json_lists_counted_columns_only(web_course, drive_root):
    assert push(web_course)[0] == 0
    cj = json.loads((mirror(drive_root, web_course) / "course.json").read_text(encoding="utf-8"))
    assert cj["format"] == 1
    assert (cj["course_id"], cj["course_name"], cj["term"]) == ("MOCK", "ฟิสิกส์จำลอง", "2027_S1")
    assert cj["folder"] == F1 and cj["prefix"] == P1
    assert cj["category_order"] == ["attendance", "homework", "midterm", "final"]
    assert cj["weights"] == {"attendance": 0.1, "homework": 0.2, "midterm": 0.3, "final": 0.4}
    assert cj["grade_boundaries"] == {"A": 80, "B+": 75, "B": 70, "C+": 65, "C": 60, "D+": 55, "D": 50}
    assert cj["rules"] == {"drop_lowest_homework": True}
    assert cj["data_mapping"] == {"attendance": ["13 Jan 2027", "20 Jan 2027", "27 Jan 2027"],
                                  "homework": ["hw1", "hw2", "hw3"], "midterm": ["midterm_exam"],
                                  "final": ["final_exam"]}
    assert cj["total_pts"] == {}
    assert cj["max_scores"] == {"hw1": 10, "hw2": 10, "hw3": 10, "midterm_exam": 30, "final_exam": 40}
    assert cj["roster"] == f"course_info/{P1}_student_info.csv"
    assert cj["files"] == sorted(f"data/{P1}_{c}.csv" for c in ("attendance", "final", "homework", "midterm"))
    att, hw, mid, fin = (f"data/{P1}_{c}.csv" for c in ("attendance", "homework", "midterm", "final"))
    assert cj["columns"] == [
        {"name": "13 Jan 2027", "header": "13 Jan 2027", "file": att, "category": "attendance"},
        {"name": "20 Jan 2027", "header": "20 Jan 2027", "file": att, "category": "attendance"},
        {"name": "27 Jan 2027", "header": "27 Jan 2027", "file": att, "category": "attendance"},
        {"name": "hw1", "header": "hw1 (10pts)", "file": hw, "category": "homework"},
        {"name": "hw2", "header": "hw2 (10pts)", "file": hw, "category": "homework"},
        {"name": "hw3", "header": "hw3 (10pts)", "file": hw, "category": "homework"},
        {"name": "midterm_exam", "header": "midterm_exam (30pts)", "file": mid, "category": "midterm"},
        {"name": "final_exam", "header": "final_exam (40pts)", "file": fin, "category": "final"},
    ]
    assert not any(col["name"].lower().startswith("total") for col in cj["columns"])


def test_edge1_2_bom_crlf_copied_bytes(tmp_path, drive_root):
    c = tmp_path / "mac" / F2
    shutil.copytree(FIXTURES / F2, c)
    code, out = push(c)
    assert code == 0, out
    m = mirror(drive_root, c)
    for p in list((c / "data").glob("*.csv")) + list((c / "course_info").glob("*.csv")):
        rel = p.relative_to(c).as_posix()
        assert (m / rel).read_bytes() == p.read_bytes(), rel
    assert (m / "data" / "9_กลศาสตร์ทดลอง_homework.csv").read_bytes().startswith(b"\xef\xbb\xbf")
    assert b"\r\n" in (m / "data" / "9_กลศาสตร์ทดลอง_homework.csv").read_bytes()


def test_edge6_duplicate_id_warning(web_course, drive_root):
    hw = web_course / "data" / f"{P1}_homework.csv"
    hw.write_text(hw.read_text(encoding="utf-8") + "9900000002,นางสาวตัวอย่าง ใจดี,10,,7,17\n", encoding="utf-8")
    code, out = push(web_course)
    assert code == 0, out
    assert f"⚠️  data/{P1}_homework.csv: Student ID 9900000002 is on 2 rows." in out


def test_edge7_missing_student_id_or_config_refused(web_course, drive_root):
    hw = web_course / "data" / f"{P1}_homework.csv"
    original = hw.read_bytes()
    hw.write_bytes(original.replace(b"Student ID,", b"Student,", 1))
    code, out = push(web_course)
    assert code == 1
    assert f"❌ data/{P1}_homework.csv has no 'Student ID' column" in out
    assert not (drive_root / "grade_courses").exists() and not (drive_root / "grade_web").exists()
    hw.write_bytes(original)
    (web_course / "course_info" / f"{P1}_config.yaml").unlink()
    code, out = push(web_course)
    assert code == 1
    assert "❌ No config in course_info/" in out
    assert not (drive_root / "grade_courses").exists() and not (drive_root / "grade_web").exists()


def test_edge8_header_only_refused(web_course, drive_root):
    fin = web_course / "data" / f"{P1}_final.csv"
    fin.write_text("Student ID,Name,final_exam (40pts),total (40pts)\n", encoding="utf-8")
    code, out = push(web_course)
    assert code == 1
    assert f"❌ data/{P1}_final.csv has no students (header only)." in out
    assert not (drive_root / "grade_courses").exists() and not (drive_root / "grade_web").exists()
    shutil.copyfile(FIXTURES / F1 / "data" / fin.name, fin)
    roster = web_course / "course_info" / f"{P1}_student_info.csv"
    roster.write_text("Student ID,Name,Class Group\n", encoding="utf-8")
    code, out = push(web_course)
    assert code == 1
    assert f"❌ course_info/{P1}_student_info.csv has no students." in out


def test_edge17_conflict_copies_refuse_push(web_course, drive_root):
    assert push(web_course)[0] == 0
    m = mirror(drive_root, web_course)
    (m / "web_edits (1).jsonl").write_text("", encoding="utf-8")
    before = tree_sha(drive_root)
    code, out = push(web_course)
    assert code == 1
    assert f"❌ Drive has duplicate files in grade_courses/mac/{F1}/: web_edits (1).jsonl." in out
    assert tree_sha(drive_root) == before


@pytest.mark.skipif(os.geteuid() == 0, reason="root can read a chmod 000 file")
def test_edge19_unreadable_journal_refuses(web_course, drive_root):
    assert push(web_course)[0] == 0
    j = put_journal(mirror(drive_root, web_course), "three_unpulled.jsonl")
    j.chmod(0)
    try:
        code, out = push(web_course, "--force")
    finally:
        j.chmod(0o644)
    assert code == 1
    assert f"❌ Couldn't read grade_courses/mac/{F1}/web_edits.jsonl (" in out
    assert not (mirror(drive_root, web_course) / "archive").exists()


def test_edge29_invalid_journal_refuses_push(web_course, drive_root):
    assert push(web_course)[0] == 0
    put_journal(mirror(drive_root, web_course), "truncated.jsonl")
    before = tree_sha(drive_root)
    for extra in ([], ["--force"]):
        code, out = push(web_course, *extra)
        assert code == 1
        assert "❌ web_edits.jsonl has lines that aren't valid JSON: 3." in out
        assert tree_sha(drive_root) == before


def test_existing_mirror_keeps_journal_and_says_so(web_course, drive_root):
    assert push(web_course)[0] == 0
    m = mirror(drive_root, web_course)
    first = json.loads((m / "push.json").read_text(encoding="utf-8"))
    code, out = push(web_course)
    assert code == 0, out
    assert "web_edits.jsonl kept (0 edits not pulled)" in out
    assert "Replacing the copy pushed" in out
    assert json.loads((m / "push.json").read_text(encoding="utf-8"))["pushed_at"] >= first["pushed_at"]


def test_debug_flag_and_one_line_error(web_course, drive_root, monkeypatch):
    import src.webpush as webpush

    def boom(*a, **k):
        raise RuntimeError("boom")
    monkeypatch.setattr(webpush, "build_course_json", boom)
    code, out, err = grader_cli(["web-push", web_course], err=True)
    assert code == 1
    assert "❌ Unexpected error: boom. Run again with --debug for the details." in out
    assert "Traceback" not in out + err
    code, out, err = grader_cli(["web-push", web_course, "--debug"], err=True)
    assert code == 1
    assert "Traceback" in out + err and "boom" in out + err
