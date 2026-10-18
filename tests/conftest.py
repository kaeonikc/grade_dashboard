import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

PREFIX = "2026_S1_Test_SEC_51"
CONFIG = """course_id: PHYS1120
course_name: Test
term: 2026_S1
sec_num: '51'
weights:
  homework: 0.3
  midterm: 0.3
  final: 0.4
data_mapping:
  homework:
  - hw1
  midterm:
  - midterm_mcq: 35
  - midterm_showsol: 5
  final:
  - final_exam
grade_boundaries:
  A: 80
  B: 70
  C: 60
  D: 50
"""


def write(path: Path, text: str, bom=False, crlf=False):
    path.parent.mkdir(parents=True, exist_ok=True)
    if crlf:
        text = text.replace("\n", "\r\n")
    data = text.encode("utf-8")
    path.write_bytes((b"\xef\xbb\xbf" if bom else b"") + data)


@pytest.fixture
def course(tmp_path) -> Path:
    """A small grade_dashboard course folder in the real on-disk format."""
    c = tmp_path / f"{PREFIX}_grading"
    write(c / "course_info" / f"{PREFIX}_config.yaml", CONFIG)
    write(c / "course_info" / f"{PREFIX}_student_info.csv",
          "Student ID,Name,Class Group\n"
          "69000001,นาย ก  หนึ่ง,G1\n69000002,นาย ข  สอง,G1\n69000003,นาย ค  สาม,G1\n")
    write(c / "data" / f"{PREFIX}_homework.csv",
          "Student ID,Name,hw1 (30pts),total (30pts)\n"
          "69000001,นาย ก  หนึ่ง,30,30.0\n69000002,นาย ข  สอง,15,15.0\n69000003,นาย ค  สาม,,0.0\n")
    write(c / "data" / f"{PREFIX}_midterm.csv",
          "Student ID,Name,midterm_mcq (35pts),midterm_showsol (5pts),total (40pts)\n"
          "69000001,นาย ก  หนึ่ง,35,5,40.0\n69000002,นาย ข  สอง,ขส,,0.0\n69000003,นาย ค  สาม,abs,2,2.0\n",
          bom=True, crlf=True)
    write(c / "data" / f"{PREFIX}_final.csv",
          "Student ID,Name,final_exam (40pts),total (40pts)\n"
          "69000001,นาย ก  หนึ่ง,40,40.0\n69000002,นาย ข  สอง,20,20.0\n69000003,นาย ค  สาม,ขส,0.0\n")
    return c


# ---- grade_web (web/SPEC.md "Test plan", layer 2). Existing fixtures above are unchanged. -------------------------
WEB_FIXTURES = ROOT / "web" / "tests" / "fixtures"
F1_NAME = "9_ฟิสิกส์จำลอง_grading"


@pytest.fixture(autouse=True)
def drive_root(tmp_path, monkeypatch) -> Path:
    """A temporary "My Drive" for every test, so nothing reaches ~/Library/CloudStorage (S1.8)."""
    root = tmp_path / "My Drive"
    root.mkdir(exist_ok=True)
    monkeypatch.setenv("GRADE_WEB_DRIVE_ROOT", str(root))
    return root


@pytest.fixture
def web_course(tmp_path) -> Path:
    """A copy of fixture F1 (web/tests/fixtures/9_ฟิสิกส์จำลอง_grading) on the Mac side."""
    import shutil
    dest = tmp_path / "mac" / F1_NAME
    shutil.copytree(WEB_FIXTURES / F1_NAME, dest)
    return dest


def build_attendance_xlsx(course: Path) -> None:
    """The attendance xlsx made from the csv codes, then a sync baseline (as `grader init` + `grader sync` leave it)."""
    import csv as _csv
    import io as _io
    import grader
    import openpyxl
    import pandas as pd
    csv_path = next((course / "data").glob("*_attendance.csv"))
    rows = list(_csv.reader(_io.StringIO(csv_path.read_bytes().decode("utf-8-sig"), newline="")))
    sessions = [h for h in rows[0][2:] if not h.lower().startswith("total")]
    reg = pd.read_csv(next((course / "course_info").glob("*_student_info.csv")), dtype=str)
    xlsx = csv_path.with_suffix(".xlsx")
    grader.save_attendance_excel(xlsx, reg, sessions)
    wb = openpyxl.load_workbook(xlsx)
    ws = wb.worksheets[0]
    head = [c.value for c in ws[1]]
    codes = {r[0]: r for r in rows[1:]}
    for r in range(2, ws.max_row + 1):
        src = codes.get(str(ws.cell(row=r, column=1).value))
        for s in sessions:
            v = src[rows[0].index(s)] if src else ""
            ws.cell(row=r, column=head.index(s) + 1).value = v or None
    wb.save(xlsx)
    grader.sync_cmd(str(course), yes=True)


@pytest.fixture
def web_course_att(web_course) -> Path:
    """F1 plus the attendance xlsx built from the csv codes and a baseline from `grader sync --yes`."""
    build_attendance_xlsx(web_course)
    return web_course
